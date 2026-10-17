# -*- coding: utf-8 -*-
"""
Excel 用例读写工具层
====================
负责与 `data/api_tables.xlsx` 交互，是「Excel 数据源」的落地实现。

相比原框架的改进：
    1. 修正原文件名拼写（execlutil -> excelutil）
    2. 工作簿带缓存：同一文件只加载一次，写入先进内存，收尾统一落盘，
       避免原实现「每条用例 load + save 整个 xlsx」的性能问题
    3. 用真实物理行号回写结果，不再依赖「行号 = 用例 id + 1」的脆弱假设
    4. 表头与列号集中映射，改表结构只需改一处

sheet 结构约定：
    test_case_list     用例表（与 Excel 表头一致）
    test_config        环境配置表
    test_result_record 结果记录表
"""
import datetime
import os

import openpyxl

from common import base
from common.exceptions import ConfigError
from config import settings
from utils.logutil import logger

# test_case_list 表头 -> 列号（1 起）
#
# 前 15 列为框架原有列；16 列起是「系统适配层」新增的可选列，
# 目的是让 Excel 数据源也能用上接口层的全部能力（否则换系统时
# 只有 YAML 用户能用 host / expects / setup 这些字段，Excel 用户用不了）。
# 新增列一律追加在末尾且可以为空，因此**旧用例表不加这些列也照常运行**。
CASE_COLUMNS = {
    "id": 1,
    "time": 2,
    "title": 3,
    "url": 4,
    "method": 5,
    "headers": 6,
    "cookies": 7,
    "request_body": 8,
    "request_type": 9,
    "relation": 10,
    "expected_code": 11,
    "response": 12,
    "ispass": 13,
    "isdel": 14,
    "tags": 15,          # 用例标签，多个用逗号分隔，如 smoke,core（用于冒烟筛选）
    # ---- 以下为可选扩展列（可整列不填） ----
    "host": 16,          # 域名别名，对应 system 档案 [hosts] 段（网关 + 多服务系统）
    "content_type": 17,  # 覆盖默认请求体类型（档案 [env] content_type 的下级覆盖）
    "expects": 18,       # 字段级断言，如 contains:body.msg=成功; gt:body.data.total=0
    "setup": 19,         # 前置步骤（造数据），JSON 数组
    "teardown": 20,      # 后置清理步骤，JSON 数组；断言失败也会执行
    "depends": 21,       # 依赖的用例 id，逗号分隔；前置未通过则本用例跳过
    "check_business": 22,  # 是否做业务码断言，默认 1；令牌/下载类接口可填 0
}

# test_config 表头 -> 列号
CONFIG_COLUMNS = {"id": 1, "web": 2, "key": 3, "value": 4}

# test_result_record 表头 -> 列号
RESULT_COLUMNS = {"id": 1, "case_id": 2, "times": 3, "response": 4, "result": 5}


class ExcelUtil:
    """Excel 读写封装（工作簿级缓存）。"""

    _cache = {}     # {文件绝对路径: openpyxl.Workbook}
    _dirty = set()  # 发生过写入的文件路径，收尾时只落盘这些，避免无谓 IO

    @classmethod
    def _mark_dirty(cls, filename):
        """标记文件已被修改，flush() 时才会真正写盘。"""
        cls._dirty.add(filename)

    def __init__(self, filename=None, sheet_name=None):
        self.filename = filename or settings.EXCEL_FILE
        self.sheet_name = sheet_name or settings.RunConfig.SHEET_NAME
        self.wb = self._load(self.filename)
        if self.sheet_name not in self.wb.sheetnames:
            raise KeyError("sheet 不存在：{}（当前可用：{}）".format(self.sheet_name, self.wb.sheetnames))
        self.sheet = self.wb[self.sheet_name]

    # ------------------------------------------------------------------
    # 内部
    # ------------------------------------------------------------------
    @classmethod
    def _load(cls, filename):
        """加载并缓存工作簿，避免重复 IO。"""
        if filename not in cls._cache:
            # 文件缺失属于「配置问题」而非「代码问题」，必须给出可操作的提示。
            # 否则 openpyxl 会抛出层层嵌套的 ZipFile 栈（FileNotFoundError ->
            # _validate_archive -> ExcelReader），让人误以为是框架坏了。
            if not os.path.isfile(filename):
                raise ConfigError(
                    "Excel 用例文件不存在：{}\n"
                    "  请检查 config/conf.ini 的 [run] 段，或用环境变量指定：\n"
                    "      AMSAPI_EXCEL_FILE=<xlsx 文件路径>\n"
                    "      AMSAPI_DATA_DIR=<数据目录>（默认读取该目录下的 api_tables.xlsx）\n"
                    "  当前生效配置：AMSAPI_EXCEL_FILE={} / AMSAPI_DATA_DIR={}".format(
                        filename,
                        os.environ.get(settings.ENV_PREFIX + "EXCEL_FILE") or "(未设置)",
                        os.environ.get(settings.ENV_PREFIX + "DATA_DIR") or "(未设置)"))
            logger.debug("加载 Excel：{}".format(filename))
            cls._cache[filename] = openpyxl.load_workbook(filename)
        return cls._cache[filename]

    def _read_rows(self, columns):
        """按列映射读取整张表，返回 dict 列表（含 __row__ 物理行号）。"""
        rows = []
        for row in range(2, self.sheet.max_row + 1):        # 第 1 行为表头
            record = {name: self.sheet.cell(row=row, column=col).value
                      for name, col in columns.items()}
            if all(value is None for value in record.values()):
                continue                                     # 跳过全空行
            record["__row__"] = row
            rows.append(record)
        return rows

    # ------------------------------------------------------------------
    # 用例表
    # ------------------------------------------------------------------
    def read_cases(self, only_runnable=True, tags=None):
        """
        读取用例列表。

        :param only_runnable: 为 True 时只返回 isdel == 1 的用例
        :param tags:          标签筛选（字符串或列表）；命中任意一个即返回。
                              为空表示不筛选。CI/CD 冒烟测试即通过此参数只取核心用例。
        """
        cases = []
        for record in self._read_rows(CASE_COLUMNS):
            cases.append(record)

        if only_runnable:
            cases = [c for c in cases if str(c.get("isdel")).strip() == "1"]

        if base.parse_tags(tags):
            cases = [c for c in cases if base.match_tags(c.get("tags"), tags)]

        logger.info("Excel[{}] 读取到 {} 条可执行用例".format(self.sheet_name, len(cases)))
        return cases

    def update_result(self, row, response, is_pass, write_response=True):
        """
        把单条用例的执行结果写回内存。

        :param row:            物理行号（来自用例的 __row__）
        :param response:       响应数据
        :param is_pass:        是否通过
        :param write_response: 是否把响应体也写入（响应体过大会拖慢写入）
        """
        now = datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        self.sheet.cell(row=row, column=CASE_COLUMNS["time"]).value = now
        if write_response:
            self.sheet.cell(row=row, column=CASE_COLUMNS["response"]).value = str(response)
        self.sheet.cell(row=row, column=CASE_COLUMNS["ispass"]).value = str(is_pass)
        self._mark_dirty(self.filename)

    # ------------------------------------------------------------------
    # 配置表 / 结果表
    # ------------------------------------------------------------------
    def read_config(self):
        """读取 test_config 表内容（本次打开的工作表）。"""
        return self._read_rows(CONFIG_COLUMNS)

    def read_results(self):
        """读取 test_result_record 表内容（本次打开的工作表）。"""
        return self._read_rows(RESULT_COLUMNS)

    def append_result(self, case_id, response, result):
        """在结果表末尾追加一条执行记录。"""
        row = self.sheet.max_row + 1
        now = datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        self.sheet.cell(row=row, column=RESULT_COLUMNS["id"]).value = row - 1
        self.sheet.cell(row=row, column=RESULT_COLUMNS["case_id"]).value = case_id
        self.sheet.cell(row=row, column=RESULT_COLUMNS["times"]).value = now
        self.sheet.cell(row=row, column=RESULT_COLUMNS["response"]).value = str(response)
        self.sheet.cell(row=row, column=RESULT_COLUMNS["result"]).value = str(result)
        self._mark_dirty(self.filename)

    # ------------------------------------------------------------------
    # 落盘
    # ------------------------------------------------------------------
    def save(self):
        """保存当前工作簿。"""
        self.wb.save(self.filename)
        logger.debug("已保存 Excel：{}".format(os.path.basename(self.filename)))

    @classmethod
    def flush(cls):
        """
        把所有「发生过写入」的工作簿统一落盘，并清空缓存（测试会话收尾调用）。

        只有被 update_result / append_result 标记过的文件才会写盘，
        因此 `--collect-only` 这类只读运行不会改动任何数据文件。
        """
        for filename in list(cls._dirty):
            wb = cls._cache.get(filename)
            if wb is not None:
                wb.save(filename)
                logger.info("结果已回写 Excel：{}".format(os.path.basename(filename)))
        cls._cache.clear()
        cls._dirty.clear()


# ----------------------------------------------------------------------
# 便捷函数（兼容原框架调用习惯）
# ----------------------------------------------------------------------
def get_sheet(filename, sheetname):
    """获取指定工作表对象（走缓存）。"""
    return ExcelUtil(filename, sheetname).sheet


def read_case_data(sheet):
    """从已打开的工作表对象读取用例（兼容旧签名）。"""
    util = ExcelUtil.__new__(ExcelUtil)          # 复用逻辑但不重复加载
    util.filename = sheet.parent.path if hasattr(sheet, "parent") else settings.EXCEL_FILE
    util.sheet_name = sheet.title
    util.wb = sheet.parent
    util.sheet = sheet
    return util.read_cases()


if __name__ == "__main__":
    for case in ExcelUtil().read_cases():
        print(case["id"], case["title"], case["url"])
