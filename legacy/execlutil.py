import json

import openpyxl
import datetime
from utils.logutil import logger
from config.settings import get_conf_path

# filename = '../data/api_tables.xlsx'
# sheetname = 'test_case_list'


def get_sheet(filename, sheetname):
    wb = openpyxl.load_workbook(filename)
    sheet = wb[sheetname]
    return sheet


def read_case_data(sheet):
    # sheet = get_sheet(filename, sheetname)
    max_raw = sheet.max_row
    case_list = []
    for i in range(2, max_raw + 1):
        if sheet.cell(row=i, column=14).value == 1:
            dict1 = dict(
                id=sheet.cell(row=i, column=1).value,
                time=sheet.cell(row=i, column=2).value,
                title=sheet.cell(row=i, column=3).value,
                url=sheet.cell(row=i, column=4).value,
                method=sheet.cell(row=i, column=5).value,
                headers=sheet.cell(row=i, column=6).value,
                cookies=sheet.cell(row=i, column=7).value,
                request_body=sheet.cell(row=i, column=8).value,
                request_type=sheet.cell(row=i, column=9).value,
                relation=sheet.cell(row=i, column=10).value,
                expected_code=sheet.cell(row=i, column=11).value,
                response=sheet.cell(row=i, column=12).value,
                ispass=sheet.cell(row=i, column=13).value,
                isdel=sheet.cell(row=i,column=14).value
            )
            case_list.append(dict1)
        else:
            continue
    return case_list


def read_config_data(sheet):
    # sheet = get_sheet(filename, sheetname)
    max_raw = sheet.max_row
    config_list = []
    for i in range(2, max_raw + 1):
        dict1 = dict(
            id=sheet.cell(row=i, column=1).value,
            web=sheet.cell(row=i, column=2).value,
            key=sheet.cell(row=i, column=3).value,
            value=sheet.cell(row=i, column=4).value
        )
        config_list.append(dict1)
    return config_list


def read_result_record(sheet):
    # sheet = get_sheet(filename, sheetname)
    max_raw = sheet.max_row
    result_list = []
    for i in range(2, max_raw + 1):
        dict1 = dict(
            id=sheet.cell(row=i, column=1).value,
            case_id=sheet.cell(row=i, column=2).value,
            time=sheet.cell(row=i, column=3).value,
            response=sheet.cell(row=i, column=4).value,
            result=sheet.cell(row=i, column=5).value
        )
        result_list.append(dict1)
    return result_list


def update_case_result(filename, sheetname, response, ispass, row_id):
    wb = openpyxl.load_workbook(filename)
    sheet = wb[sheetname]
    current_time = datetime.datetime.now().strftime('%Y-%m-%d %H:%M:%S')
    sheet.cell(row_id, 2).value = current_time
    # sheet.cell(row_id, 12).value = {json.dumps(response, ensure_ascii=False)}
    sheet.cell(row_id, 12).value = str(response)
    sheet.cell(row_id, 13).value = str(ispass)
    wb.save(filename)


if __name__ == '__main__':
    filename = '../data/api_tables.xlsx'
    sheetname = 'test_case_list'
    get_sheet('../data/api_tables.xlsx', 'test_case_list')
    case = read_case_data(get_sheet('../data/api_tables.xlsx', 'test_case_list'))
#     d_time = datetime.datetime.now().strftime('%Y-%m-%d %H:%M:%S')
#     print(type(d_time))
    print(case)
#     print(eval(case))
#     print(type(case))
#     res_data = {'code': 200, 'body': {'status': 500, 'message': '用户名不能为空'}, 'cookies': {}}
#     print(type(str(res_data)))
#     update_case_result('../data/api_tables.xlsx', 'test', response=str(res_data), ispass='false', row_id=1)
