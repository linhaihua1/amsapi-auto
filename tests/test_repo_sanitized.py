# -*- coding: utf-8 -*-
"""
单元测试：仓库开源合规（不允许残留真实环境信息）
=================================================
本文件守护的是一条**约束**而不是一段功能：

    仓库会被公开分发，因此任何随仓库发布的文件里都不允许出现
    真实内网地址、真实账号、真实口令、真实业务标识、真实用户隐私。

为什么值得写成测试：
    这类信息往往藏在注释、示例数据、TODO、历史归档里，
    靠人工 review 一定会漏。写成断言后，只要有人误提交就会立刻红。

扫描范围：仓库根下所有文本文件 + Excel 用例文件。
排除：venv / .git / __pycache__ / report / log / .idea（构建与本地产物）。
"""
import os

import openpyxl
import pytest

from config import settings

ROOT = settings.BASE_DIR

# 不参与扫描的目录（本地产物 / 第三方 / 版本控制）
SKIP_DIRS = {"venv", ".venv", ".git", "__pycache__", ".idea", "report", "log",
             "node_modules", ".pytest_cache", "allure-results", "allure-report"}

# 文本类后缀（其余文件按二进制跳过）
TEXT_SUFFIX = (".py", ".md", ".ini", ".yaml", ".yml", ".json", ".txt",
               ".cfg", ".toml", ".properties", ".jenkinsfile", "")

# 黑名单：(关键字, 说明)。全部是「只在真实环境里才会出现」的字面量。
#
# 注意：关键字刻意用**片段拼接**而不是整串字面量 —— 否则本文件自己就会命中
# 自己的扫描规则（守卫代码不能成为违规样本）。
def _kw(*parts):
    return "".join(parts)


FORBIDDEN = [
    (_kw("139", ".186."), "真实内网 IP 前缀"),
    (_kw("linhai", "hua"), "真实测试账号"),
    (_kw("dangjian", "yun"), "真实系统/项目代号"),
    (_kw("cpms", "-"), "真实应用编码前缀"),
    (_kw("material", "-types"), "真实业务数据标识"),
    (_kw("2", "chu"), "真实组织编码"),
    (_kw("1317249", "4100"), "真实手机号"),
    (_kw("03d75d44", "cb444543d641283fb2adb5ce"), "真实口令摘要"),
    (_kw("aac89144", "300e665794460f1021d556ed"), "真实签名值"),
]


def _iter_files():
    """遍历需要扫描的文件（跳过产物目录）。"""
    for base, dirs, names in os.walk(ROOT):
        dirs[:] = [d for d in dirs if d not in SKIP_DIRS]
        for name in names:
            yield os.path.join(base, name)


def _text_files():
    """只取文本类文件。"""
    for path in _iter_files():
        if os.path.splitext(path)[1].lower() in TEXT_SUFFIX:
            yield path


def _rel(path):
    return os.path.relpath(path, ROOT)


class TestNoRealEnvInfoInRepo:
    """仓库里的文本文件不得包含真实环境信息。"""

    @pytest.mark.parametrize("keyword,label", FORBIDDEN, ids=[f[1] for f in FORBIDDEN])
    def test_text_files(self, keyword, label):
        hits = []
        for path in _text_files():
            try:
                text = open(path, encoding="utf-8").read()
            except (UnicodeDecodeError, OSError):
                continue
            if keyword in text:
                hits.append(_rel(path))
        assert not hits, "以下文件残留{}（{}）：\n  {}".format(
            label, keyword, "\n  ".join(sorted(hits)))

    def test_excel_case_file(self):
        """Excel 用例文件（二进制）单独扫：里面有真实账号最多。"""
        wb = openpyxl.load_workbook(settings.EXCEL_FILE, data_only=True)
        cells = []
        for sheet in wb.worksheets:
            for row in sheet.iter_rows(values_only=True):
                for value in row:
                    if value is not None:
                        cells.append(str(value))
        blob = "\n".join(cells)
        for keyword, label in FORBIDDEN:
            assert keyword not in blob, "Excel 中残留{}（{}）".format(label, keyword)


class TestShippedSampleData:
    """随仓库发布的示例数据必须能「离线开箱即跑」。"""

    def test_excel_has_runnable_cases(self):
        wb = openpyxl.load_workbook(settings.EXCEL_FILE, data_only=True)
        ws = wb[settings.RunConfig.SHEET_NAME]
        rows = [row for row in ws.iter_rows(min_row=2, values_only=True)
                if row and str(row[13]).strip() == "1"]
        assert len(rows) >= 3, "示例 Excel 至少要留几条可执行用例"

    def test_excel_urls_are_relative(self):
        """
        示例用例一律用相对路径。

        理由：绝对地址会把某台机器的环境钉死在仓库里，开源后别人跑不通；
        相对路径 + AMSAPI_BASE_URL 才是可移植的写法。
        """
        wb = openpyxl.load_workbook(settings.EXCEL_FILE, data_only=True)
        ws = wb[settings.RunConfig.SHEET_NAME]
        for index, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            if not row or row[3] is None:
                continue
            url = str(row[3])
            assert not url.lower().startswith(("http://", "https://")), \
                "Excel 第 {} 行的 url 应改为相对路径：{}".format(index, url)

    def test_yaml_cases_use_relative_paths(self):
        from utils.yamlutil import read_cases
        for case in read_cases():
            url = str(case.get("url") or "")
            assert not url.lower().startswith(("http://", "https://")), \
                "YAML 用例 {} 的 url 应改为相对路径：{}".format(case.get("title"), url)


class TestSanitizedConfig:
    """配置文件里不得预置真实凭据。"""

    def test_db_password_not_prefilled(self):
        assert settings.DB_CONFIG["password"] == ""

    def test_conf_ini_has_no_password_value(self):
        text = open(settings.CONF_FILE, encoding="utf-8").read()
        for line in text.splitlines():
            stripped = line.strip()
            if stripped.startswith("password") and "=" in stripped:
                assert stripped.split("=", 1)[1].strip() == "", \
                    "conf.ini 的 password 不应预置任何值：{}".format(stripped)
