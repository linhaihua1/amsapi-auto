# -*- coding: utf-8 -*-
"""
单元测试：utils/excelutil.py（Excel 用例读写层）
===============================================
关注三类行为：
    1. 读取与筛选（isdel / tags / 跳过空行 / 物理行号）
    2. 结果回写（写内存 → 脏标记 → flush 才落盘）
    3. 只读运行不落盘（`--collect-only` 无副作用的保障）

全部在 tmp_path 的副本上进行，绝不触碰 data/api_tables.xlsx。
"""
import os
import shutil

import openpyxl
import pytest

from config import settings
from utils.excelutil import (CASE_COLUMNS, CONFIG_COLUMNS, RESULT_COLUMNS,
                             ExcelUtil, get_sheet, read_case_data)

SHEET = "test_case_list"


def _md5(path):
    import hashlib
    with open(path, "rb") as fp:
        return hashlib.md5(fp.read()).hexdigest()


# ======================================================================
# 一、列映射自检
# ======================================================================
class TestColumnMapping:
    def test_case_columns_are_continuous(self):
        """列号必须连续且从 1 开始，否则读写会错位。"""
        assert sorted(CASE_COLUMNS.values()) == list(range(1, len(CASE_COLUMNS) + 1))

    def test_real_excel_header_matches_mapping(self):
        """
        真实 Excel 的表头必须与代码中的列映射**连续前缀一致**。

        为什么是「前缀」而不是「完全相等」：
            前 15 列是框架原有列，必须严格对齐，否则读写会整体错位；
            第 16 列起是系统适配层新增的可选列（host / content_type / expects /
            setup / teardown / depends / check_business），允许用户在自己的用例表里
            按需追加——没写的列读出来是 None，不会影响既有用例。

        这里从第 1 列往后读，遇到第一个空表头即停止，因此「文件里连续定义的列」
        仍必须与映射逐位同名，防止有人在中间插一列导致静默错位。
        """
        ws = openpyxl.load_workbook(settings.EXCEL_FILE)[settings.RunConfig.SHEET_NAME]
        expected = sorted(CASE_COLUMNS, key=CASE_COLUMNS.get)

        header = []
        for col in range(1, len(expected) + 1):
            value = ws.cell(1, col).value
            if value is None:
                break
            header.append(value)

        assert header == expected[:len(header)]

    def test_config_and_result_columns(self):
        assert CONFIG_COLUMNS == {"id": 1, "web": 2, "key": 3, "value": 4}
        assert RESULT_COLUMNS == {"id": 1, "case_id": 2, "times": 3,
                                  "response": 4, "result": 5}


# ======================================================================
# 二、读取与筛选
# ======================================================================
class TestReadCases:
    def test_reads_only_runnable(self, excel_file):
        cases = ExcelUtil(excel_file, SHEET).read_cases()
        assert [c["id"] for c in cases] == [1, 3, 4]           # isdel=0 的被排除

    def test_reads_optional_adapter_columns(self, tmp_path):
        """第 16 列起的可选适配列也要能读出来，Excel 数据源才能用上系统适配能力。"""
        path = tmp_path / "cases_opt.xlsx"
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = SHEET
        ws.append(sorted(CASE_COLUMNS, key=CASE_COLUMNS.get))
        ws.append([1, None, "用例", "/x", "get", None, None, None, None, "None", 200,
                   None, None, 1, "smoke", "order", "application/json",
                   "gt:body.n=0", '[{"url": "/a"}]', '[{"url": "/b"}]', "1", "0"])
        wb.save(str(path))

        case = ExcelUtil(str(path), SHEET).read_cases()[0]
        assert case["host"] == "order"
        assert case["content_type"] == "application/json"
        assert case["expects"] == "gt:body.n=0"
        assert case["setup"] == '[{"url": "/a"}]'
        assert case["teardown"] == '[{"url": "/b"}]'
        assert case["depends"] == "1"
        assert case["check_business"] == "0"

    def test_legacy_15_column_file_still_reads(self, excel_file):
        """旧用例表只写到 tags 列（15 列），新增列缺失时应读出 None 而不是报错。"""
        case = ExcelUtil(excel_file, SHEET).read_cases()[0]
        assert case["host"] is None
        assert case["expects"] is None
        assert case["depends"] is None
        assert case["check_business"] is None

    def test_skips_blank_row(self, excel_file):
        """全空行不应被当成用例。"""
        cases = ExcelUtil(excel_file, SHEET).read_cases(only_runnable=False)
        assert len(cases) == 4                                 # 共 4 条有内容的行

    def test_physical_row_number(self, excel_file):
        """__row__ 必须是真实物理行号，供结果回写定位。"""
        cases = ExcelUtil(excel_file, SHEET).read_cases(only_runnable=False)
        assert [c["__row__"] for c in cases] == [2, 3, 5, 6]

    def test_all_cases_when_only_runnable_false(self, excel_file):
        cases = ExcelUtil(excel_file, SHEET).read_cases(only_runnable=False)
        assert [c["id"] for c in cases] == [1, 2, 3, 4]

    def test_filter_by_single_tag(self, excel_file):
        cases = ExcelUtil(excel_file, SHEET).read_cases(tags="smoke")
        assert [c["id"] for c in cases] == [1, 3]

    def test_filter_by_multi_tag(self, excel_file):
        cases = ExcelUtil(excel_file, SHEET).read_cases(tags="smoke,regression")
        assert [c["id"] for c in cases] == [1, 3, 4]

    def test_filter_tag_not_matched_returns_empty(self, excel_file):
        assert ExcelUtil(excel_file, SHEET).read_cases(tags="not-exist") == []

    def test_no_tag_filter_returns_all(self, excel_file):
        assert len(ExcelUtil(excel_file, SHEET).read_cases(tags=None)) == 3

    def test_missing_sheet_raises_keyerror(self, excel_file):
        with pytest.raises(KeyError):
            ExcelUtil(excel_file, "no_such_sheet")

    def test_excel_content_not_evaluated(self, excel_file):
        """读取到的内容保持原始字符串，不做任何代码求值。"""
        case = ExcelUtil(excel_file, SHEET).read_cases()[0]
        assert case["headers"] == '{"Content-Type": "application/json"}'
        assert case["relation"] == "ticket=body.data.ticket"


# ======================================================================
# 三、结果回写与落盘
# ======================================================================
class TestWriteBack:
    def test_update_result_writes_memory_only(self, excel_file):
        """update_result 只改内存，未 flush 前文件内容不应变化。"""
        before = _md5(excel_file)
        util = ExcelUtil(excel_file, SHEET)
        util.update_result(2, {"code": 200}, True)
        assert _md5(excel_file) == before
        assert util.sheet.cell(2, CASE_COLUMNS["ispass"]).value == "True"

    def test_flush_persists_dirty_file(self, excel_file):
        before = _md5(excel_file)
        util = ExcelUtil(excel_file, SHEET)
        util.update_result(2, {"code": 200, "body": {"status": 200}}, True)
        ExcelUtil.flush()

        assert _md5(excel_file) != before
        ws = openpyxl.load_workbook(excel_file)[SHEET]
        assert ws.cell(2, CASE_COLUMNS["ispass"]).value == "True"
        assert ws.cell(2, CASE_COLUMNS["time"]).value          # 时间已写入
        assert "status" in ws.cell(2, CASE_COLUMNS["response"]).value

    def test_flush_is_idempotent_after_clean(self, excel_file):
        """无写入时 flush 不应改动文件。"""
        before = _md5(excel_file)
        ExcelUtil(excel_file, SHEET).read_cases()
        ExcelUtil.flush()
        assert _md5(excel_file) == before

    def test_readonly_run_does_not_touch_file(self, excel_file):
        """只读运行（模拟 --collect-only）不得修改任何文件。"""
        before = _md5(excel_file)
        ExcelUtil(excel_file, SHEET).read_cases()
        ExcelUtil(excel_file, SHEET).read_cases(tags="smoke")
        ExcelUtil.flush()
        assert _md5(excel_file) == before

    def test_failed_case_written_as_false(self, excel_file):
        util = ExcelUtil(excel_file, SHEET)
        util.update_result(2, {"code": 500}, False)
        ExcelUtil.flush()
        ws = openpyxl.load_workbook(excel_file)[SHEET]
        assert ws.cell(2, CASE_COLUMNS["ispass"]).value == "False"

    def test_write_response_can_be_disabled(self, excel_file):
        util = ExcelUtil(excel_file, SHEET)
        util.update_result(2, {"code": 200}, True, write_response=False)
        ExcelUtil.flush()
        ws = openpyxl.load_workbook(excel_file)[SHEET]
        assert ws.cell(2, CASE_COLUMNS["response"]).value is None

    def test_dirty_is_per_file(self, excel_file, tmp_path):
        """只落盘被改动的文件，未改动的文件不得被重写。"""
        other = tmp_path / "other.xlsx"
        shutil.copy(excel_file, str(other))
        other_md5 = _md5(str(other))

        util_a = ExcelUtil(excel_file, SHEET)
        ExcelUtil(str(other), SHEET)                      # 只加载，不改
        util_a.update_result(5, {}, True)
        ExcelUtil.flush()

        assert _md5(str(other)) == other_md5
        assert _md5(excel_file) != other_md5

    def test_append_result(self, excel_file):
        util = ExcelUtil(excel_file, "test_result_record")
        util.append_result(case_id=99, response={"status": 200}, result=True)
        ExcelUtil.flush()

        ws = openpyxl.load_workbook(excel_file)["test_result_record"]
        last = ws.max_row
        assert ws.cell(last, RESULT_COLUMNS["case_id"]).value == 99
        assert ws.cell(last, RESULT_COLUMNS["result"]).value == "True"
        assert ws.cell(last, RESULT_COLUMNS["times"]).value

    def test_save_directly(self, excel_file):
        util = ExcelUtil(excel_file, SHEET)
        util.sheet.cell(2, CASE_COLUMNS["ispass"]).value = "MANUAL"
        util.save()
        ws = openpyxl.load_workbook(excel_file)[SHEET]
        assert ws.cell(2, CASE_COLUMNS["ispass"]).value == "MANUAL"


# ======================================================================
# 四、缓存与配置表读取
# ======================================================================
class TestCacheAndOtherSheets:
    def test_workbook_cached(self, excel_file):
        a = ExcelUtil(excel_file, SHEET)
        b = ExcelUtil(excel_file, SHEET)
        assert a.wb is b.wb                             # 同一文件只加载一次

    def test_read_config_sheet(self, excel_file):
        rows = ExcelUtil(excel_file, "test_config").read_config()
        assert len(rows) == 2
        assert rows[0]["web"] == "gov"
        assert rows[0]["value"] == "http://192.0.2.10"

    def test_read_result_sheet(self, excel_file):
        rows = ExcelUtil(excel_file, "test_result_record").read_results()
        assert len(rows) == 1
        assert rows[0]["case_id"] == "1"

    def test_get_sheet_helper(self, excel_file):
        assert get_sheet(excel_file, SHEET).title == SHEET

    def test_read_case_data_helper(self, excel_file):
        """兼容旧签名的读取函数应能正常工作。"""
        cases = read_case_data(get_sheet(excel_file, SHEET))
        assert [c["id"] for c in cases] == [1, 3, 4]

    def test_default_filename_from_settings(self):
        assert ExcelUtil().filename == settings.EXCEL_FILE

    def test_default_sheet_from_settings(self):
        assert ExcelUtil(settings.EXCEL_FILE).sheet_name == settings.RunConfig.SHEET_NAME
